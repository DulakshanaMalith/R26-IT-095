from pathlib import Path
from typing import List
from openpyxl import load_workbook
from app.models.cohort_import import ParsedRow

STUDENT_SHEET_NAME = "Students"
STUDENT_HEADERS = [
    "StudentID",
    "React",
    "HTML_CSS",
    "Angular",
    "Vue",
    "NodeJS",
    "Express",
    "Java",
    "PHP",
    "FastAPI",
    "MongoDB",
    "MySQL",
    "PostgreSQL",
    "Firebase",
    "Python",
    "TensorFlow",
    "Pandas",
]

PROJECT_SHEET_NAME = "Projects"
PROJECT_HEADERS = [
    "ProjectID",
    "ProjectTitle",
    "TeamSize",
    "Status",
]

PROJECT_REQUIREMENTS_SHEET_NAME = "ProjectRequirements"
PROJECT_REQUIREMENT_HEADERS = [
    "ProjectID",
    "Technology",
    "MinLevel",
    "RequiredMembers",
]

PREFERENCES_SHEET_NAME = "Preferences"
PREFERENCE_HEADERS = [
    "StudentID",
    "Preference1",
    "Preference2",
    "Preference3",
]

DOMAINS_SHEET_NAME = "Domains"
DOMAIN_HEADERS = [
    "Domain",
]

PROJECT_DOMAINS_SHEET_NAME = "ProjectDomains"
PROJECT_DOMAIN_HEADERS = [
    "ProjectID",
    "Domain",
]

SUPERVISORS_SHEET_NAME = "Supervisors"
SUPERVISOR_HEADERS = [
    "SupervisorID",
    "SupervisorName",
    "MaximumTeams",
    "CurrentLoad",
]

SUPERVISOR_DOMAINS_SHEET_NAME = "SupervisorDomains"
SUPERVISOR_DOMAIN_HEADERS = [
    "SupervisorID",
    "Type",
    "Domain",
]

def normalize_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip()

def is_empty_value(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False

def _parse_sheet(
    file_path: str | Path,
    sheet_name: str,
    required_headers: List[str],
) -> List[ParsedRow]:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Required sheet '{sheet_name}' was not found."
            )

        worksheet = workbook[sheet_name]

        try:
            header_row = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                )
            )
        except StopIteration:
            raise ValueError(
                f"Sheet '{sheet_name}' is empty."
            )

        headers = [
            normalize_header(value)
            for value in header_row
        ]

        missing_headers = [
            header
            for header in required_headers
            if header not in headers
        ]

        if missing_headers:
            raise ValueError(
                f"{sheet_name} sheet is missing required columns: "
                + ", ".join(missing_headers)
            )

        header_indexes = {
            header: headers.index(header)
            for header in required_headers
        }

        parsed_rows: List[ParsedRow] = []

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            values = {}

            for header in required_headers:
                column_index = header_indexes[header]
                values[header] = (
                    row[column_index]
                    if column_index < len(row)
                    else None
                )

            if all(
                is_empty_value(value)
                for value in values.values()
            ):
                continue

            parsed_rows.append(
                ParsedRow(
                    sheet=sheet_name,
                    row_number=row_number,
                    values=values,
                )
            )

        return parsed_rows
    finally:
        workbook.close()

def parse_students_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=STUDENT_SHEET_NAME,
        required_headers=STUDENT_HEADERS,
    )

def parse_projects_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=PROJECT_SHEET_NAME,
        required_headers=PROJECT_HEADERS,
    )

def parse_project_requirements_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=PROJECT_REQUIREMENTS_SHEET_NAME,
        required_headers=PROJECT_REQUIREMENT_HEADERS,
    )

def parse_preferences_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=PREFERENCES_SHEET_NAME,
        required_headers=PREFERENCE_HEADERS,
    )

def parse_domains_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=DOMAINS_SHEET_NAME,
        required_headers=DOMAIN_HEADERS,
    )

def parse_project_domains_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=PROJECT_DOMAINS_SHEET_NAME,
        required_headers=PROJECT_DOMAIN_HEADERS,
    )

def parse_supervisors_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=SUPERVISORS_SHEET_NAME,
        required_headers=SUPERVISOR_HEADERS,
    )

def parse_supervisor_domains_sheet(
    file_path: str | Path,
) -> List[ParsedRow]:
    return _parse_sheet(
        file_path=file_path,
        sheet_name=SUPERVISOR_DOMAINS_SHEET_NAME,
        required_headers=SUPERVISOR_DOMAIN_HEADERS,
    )