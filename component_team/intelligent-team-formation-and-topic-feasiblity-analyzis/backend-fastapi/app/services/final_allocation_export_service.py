from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

def _percent(value: float) -> str:
    return f"{float(value or 0) * 100:.1f}%"

def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1E293B")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 10), 45)

def build_excel_export(allocation: dict) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary_rows = [
        ("Allocation ID", allocation["allocation_id"]),
        ("Created At", allocation["created_at"]),
        ("Status", allocation.get("status") or ""),
        ("Revision Number", allocation.get("revision_number", 1)),
        ("Allocation Source", allocation.get("allocation_source") or ""),
        ("Parent Allocation ID", allocation.get("parent_allocation_id") or ""),
        ("Change Reason", allocation.get("change_reason") or ""),
        ("Source Workbook", allocation.get("source_file_name") or ""),
        ("Algorithm", allocation["algorithm"]),
        ("Optimizer Version", allocation["optimizer_version"]),
        ("Selected Solution", allocation["solution_id"]),
        ("Solution Role", allocation["solution_role"]),
        ("Students", allocation["student_count"]),
        ("Projects / Teams", allocation["project_count"]),
        ("Students per Team", allocation["students_per_team"]),
        ("Technical Coverage", _percent(allocation["technical_requirement_coverage"])),
        ("Technical Deficit", _percent(allocation["technical_requirement_deficit"])),
        ("Preference Satisfaction", _percent(allocation["preference_satisfaction"])),
        ("Preference Dissatisfaction", _percent(allocation["preference_dissatisfaction"])),
        ("Integrity Valid", allocation["integrity_valid"]),
    ]
    for row in summary_rows:
        summary.append(list(row))
    _style_sheet(summary)
    teams_sheet = workbook.create_sheet("Teams")
    teams_sheet.append(["Team", "Project ID", "Project Title", "Team Size", "Students", "Technical Coverage", "Technical Deficit", "Preference Satisfaction", "Supervisor", "Supervisor ID", "Match Status"])
    members_sheet = workbook.create_sheet("TeamMembers")
    members_sheet.append(["Team", "Project ID", "Student ID", "Preference Rank", "Dissatisfaction", "Ranked Projects", "Relevant Skills"])
    requirements_sheet = workbook.create_sheet("TechnicalRequirements")
    requirements_sheet.append(["Team", "Project ID", "Technology", "Min Level", "Required Members", "Qualified Members", "Qualified Students", "Coverage", "Status"])
    supervisors_sheet = workbook.create_sheet("SupervisorAssignments")
    supervisors_sheet.append(["Team", "Project ID", "Project Title", "Supervisor ID", "Supervisor Name", "Match Status", "Project Domains", "Expertise Matches", "Interest Matches", "Current Load", "Maximum Teams", "Projected Load", "Remaining Capacity", "Explanation"])
    for team in allocation["teams"]:
        supervisor = team.get("supervisor") or {}
        preference_satisfaction = 1.0 - float(team["preference_summary"]["average_dissatisfaction"])
        student_ids = [student["student_id"] for student in team["students"]]
        teams_sheet.append([
            f"T{team['team_number']:03d}", team["project_id"], team["project_title"], team["team_size"], ", ".join(student_ids),
            _percent(team["technical_coverage"]), _percent(team["technical_deficit"]), _percent(preference_satisfaction),
            supervisor.get("supervisor_name", ""), supervisor.get("supervisor_id", ""), supervisor.get("match_status", ""),
        ])
        for student in team["students"]:
            skills = ", ".join(f"{technology}:{level}" for technology, level in sorted((student.get("relevant_skills") or {}).items()))
            members_sheet.append([
                f"T{team['team_number']:03d}", team["project_id"], student["student_id"], student.get("preference_rank"),
                student.get("dissatisfaction", 0.0), ", ".join(student.get("ranked_projects") or []), skills,
            ])
        for requirement in team["requirements"]:
            requirements_sheet.append([
                f"T{team['team_number']:03d}", team["project_id"], requirement["technology"], requirement["min_level"],
                requirement["required_members"], requirement["qualified_members"], ", ".join(requirement.get("qualified_students") or []),
                _percent(requirement["coverage"]), requirement["status"],
            ])
        supervisors_sheet.append([
            f"T{team['team_number']:03d}", team["project_id"], team["project_title"], supervisor.get("supervisor_id", ""),
            supervisor.get("supervisor_name", ""), supervisor.get("match_status", ""), ", ".join(supervisor.get("project_domains") or []),
            ", ".join(supervisor.get("expertise_matches") or []), ", ".join(supervisor.get("interest_matches") or []),
            supervisor.get("current_load", ""), supervisor.get("maximum_teams", ""), supervisor.get("projected_load", ""),
            supervisor.get("remaining_capacity", ""), supervisor.get("explanation", ""),
        ])
    for sheet in [teams_sheet, members_sheet, requirements_sheet, supervisors_sheet]:
        _style_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

def build_pdf_export(allocation: dict) -> BytesIO:
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("AllocationTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=18, leading=22)
    body_style = ParagraphStyle("AllocationBody", parent=styles["BodyText"], fontSize=8, leading=10)
    header_style = ParagraphStyle("AllocationHeader", parent=body_style, textColor=colors.white, fontName="Helvetica-Bold")
    source_label = "Staff-revised allocation" if allocation.get("allocation_source") == "MANUAL_REVISION" else "Optimizer-selected allocation"
    story = [
        Paragraph("Final Team Allocation Report", title_style),
        Spacer(1, 4 * mm),
        Paragraph(
            f"Allocation ID: <b>{allocation['allocation_id']}</b> | Status: <b>{allocation.get('status', '')}</b> | Revision: <b>{allocation.get('revision_number', 1)}</b> | Source: <b>{source_label}</b>",
            styles["BodyText"],
        ),
        Paragraph(
            f"Students: <b>{allocation['student_count']}</b> | Teams: <b>{allocation['project_count']}</b> | Target team size: <b>{allocation['students_per_team']}</b> | Technical coverage: <b>{_percent(allocation['technical_requirement_coverage'])}</b> | Preference satisfaction: <b>{_percent(allocation['preference_satisfaction'])}</b>",
            styles["BodyText"],
        ),
    ]
    if allocation.get("parent_allocation_id"):
        story.append(Paragraph(f"Parent allocation: <b>{allocation['parent_allocation_id']}</b> | Change reason: {allocation.get('change_reason') or ''}", styles["BodyText"]))
    story.append(Spacer(1, 5 * mm))
    table_data = [[
        Paragraph("Team", header_style), Paragraph("Project", header_style), Paragraph("Students", header_style),
        Paragraph("Supervisor", header_style), Paragraph("Match", header_style), Paragraph("Tech Coverage", header_style),
        Paragraph("Preference Satisfaction", header_style),
    ]]
    for team in allocation["teams"]:
        supervisor = team.get("supervisor") or {}
        preference_satisfaction = 1.0 - float(team["preference_summary"]["average_dissatisfaction"])
        table_data.append([
            Paragraph(f"T{team['team_number']:03d}", body_style),
            Paragraph(f"<b>{team['project_id']}</b><br/>{team['project_title']}", body_style),
            Paragraph(", ".join(student["student_id"] for student in team["students"]), body_style),
            Paragraph(f"{supervisor.get('supervisor_name', '')}<br/>{supervisor.get('supervisor_id', '')}", body_style),
            Paragraph(supervisor.get("match_status", ""), body_style),
            Paragraph(_percent(team["technical_coverage"]), body_style),
            Paragraph(_percent(preference_satisfaction), body_style),
        ])
    table = Table(table_data, repeatRows=1, colWidths=[16 * mm, 62 * mm, 62 * mm, 42 * mm, 35 * mm, 28 * mm, 32 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output
