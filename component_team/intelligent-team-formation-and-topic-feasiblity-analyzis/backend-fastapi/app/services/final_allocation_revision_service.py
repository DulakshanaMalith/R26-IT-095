from collections import Counter, defaultdict
from io import BytesIO
from math import ceil
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.db_models.final_allocation import FinalAllocation
from app.services.final_allocation_service import get_final_allocation, persist_final_allocation

REQUIRED_REVISION_SHEETS = {"AllocationInfo", "TeamAssignments", "SupervisorAssignments"}

def _normalize_id(value) -> str:
    return "" if value is None else str(value).strip().upper()

def _as_int(value, field: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field} must be an integer.")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field} must be an integer.")
    return number

def _headers(sheet) -> dict[str, int]:
    result = {}
    for index, cell in enumerate(sheet[1], start=1):
        if cell.value is not None:
            result[str(cell.value).strip()] = index
    return result

def _require_headers(sheet, required: set[str]) -> dict[str, int]:
    headers = _headers(sheet)
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Sheet '{sheet.title}' is missing required column(s): {', '.join(missing)}.")
    return headers

def _style_table_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1E293B")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 44)

def build_editable_allocation_workbook(allocation: dict, reference_data: dict | None) -> BytesIO:
    if not reference_data:
        raise ValueError("This allocation does not contain the original cohort reference data required for safe revision. Save a new allocation with the source workbook or attach the original source workbook first.")
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    instructions = [
        ["Editable Allocation Revision Workbook"],
        ["Purpose", "Make controlled manual changes to team membership/project assignment and/or supervisor assignment."],
        ["Editable sheet", "TeamAssignments: edit TeamNumber and ProjectID. Do not add/remove/rename StudentID values."],
        ["Editable sheet", "SupervisorAssignments: edit SupervisorID. TeamNumber and ProjectID must match TeamAssignments."],
        ["Do not edit", "AllocationInfo, ReferenceProjects, ReferenceSupervisors."],
        ["Important", "Uploading this workbook does not overwrite history. After validation and confirmation, the current allocation is archived and a new ACTIVE revision is created."],
        ["Important", "Technical coverage and preference metrics are recalculated by the backend from stored source data. Values typed into this workbook are never trusted as calculated results."],
    ]
    for row in instructions:
        readme.append(row)
    readme["A1"].font = Font(bold=True, size=16, color="1E3A8A")
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 110
    readme["B2"].alignment = Alignment(wrap_text=True)
    for row in range(2, readme.max_row + 1):
        readme.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    info = workbook.create_sheet("AllocationInfo")
    info.append(["Field", "Value"])
    info_rows = [
        ("ParentAllocationID", allocation["allocation_id"]),
        ("CurrentRevision", allocation.get("revision_number", 1)),
        ("CurrentStatus", allocation.get("status", "")),
        ("AllocationSource", allocation.get("allocation_source", "")),
        ("StudentsPerTeam", allocation["students_per_team"]),
        ("StudentCount", allocation["student_count"]),
        ("TeamCount", allocation["project_count"]),
    ]
    for row in info_rows:
        info.append(list(row))
    _style_table_sheet(info)
    team_sheet = workbook.create_sheet("TeamAssignments")
    team_sheet.append(["StudentID", "TeamNumber", "ProjectID"])
    for team in allocation["teams"]:
        for student in sorted(team["students"], key=lambda item: item["student_id"]):
            team_sheet.append([student["student_id"], team["team_number"], team["project_id"]])
    _style_table_sheet(team_sheet)
    supervisor_sheet = workbook.create_sheet("SupervisorAssignments")
    supervisor_sheet.append(["TeamNumber", "ProjectID", "SupervisorID"])
    for team in allocation["teams"]:
        supervisor = team.get("supervisor") or {}
        supervisor_sheet.append([team["team_number"], team["project_id"], supervisor.get("supervisor_id", "")])
    _style_table_sheet(supervisor_sheet)
    projects_sheet = workbook.create_sheet("ReferenceProjects")
    projects_sheet.append(["ProjectID", "ProjectTitle"])
    for project in sorted(reference_data.get("projects") or [], key=lambda item: str(item.get("project_id") or "")):
        projects_sheet.append([project.get("project_id"), project.get("project_title")])
    _style_table_sheet(projects_sheet)
    supervisors_sheet = workbook.create_sheet("ReferenceSupervisors")
    supervisors_sheet.append(["SupervisorID", "SupervisorName", "MaximumTeams", "CurrentLoad"])
    for supervisor in sorted(reference_data.get("supervisors") or [], key=lambda item: str(item.get("supervisor_id") or "")):
        supervisors_sheet.append([supervisor.get("supervisor_id"), supervisor.get("supervisor_name"), supervisor.get("maximum_teams"), supervisor.get("current_load")])
    _style_table_sheet(supervisors_sheet)
    if projects_sheet.max_row >= 2:
        project_validation = DataValidation(type="list", formula1=f"'ReferenceProjects'!$A$2:$A${projects_sheet.max_row}", allow_blank=False)
        team_sheet.add_data_validation(project_validation)
        project_validation.add(f"C2:C{max(team_sheet.max_row, 2)}")
    if supervisors_sheet.max_row >= 2:
        supervisor_validation = DataValidation(type="list", formula1=f"'ReferenceSupervisors'!$A$2:$A${supervisors_sheet.max_row}", allow_blank=False)
        supervisor_sheet.add_data_validation(supervisor_validation)
        supervisor_validation.add(f"C2:C{max(supervisor_sheet.max_row, 2)}")
    max_team = allocation["project_count"]
    team_validation = DataValidation(type="whole", operator="between", formula1="1", formula2=str(max_team), allow_blank=False)
    team_sheet.add_data_validation(team_validation)
    team_validation.add(f"B2:B{max(team_sheet.max_row, 2)}")
    editable_fill = PatternFill("solid", fgColor="FEF3C7")
    locked_fill = PatternFill("solid", fgColor="E2E8F0")
    for row in range(2, team_sheet.max_row + 1):
        team_sheet.cell(row, 1).fill = locked_fill
        team_sheet.cell(row, 2).fill = editable_fill
        team_sheet.cell(row, 3).fill = editable_fill
    for row in range(2, supervisor_sheet.max_row + 1):
        supervisor_sheet.cell(row, 1).fill = locked_fill
        supervisor_sheet.cell(row, 2).fill = locked_fill
        supervisor_sheet.cell(row, 3).fill = editable_fill
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

def _read_allocation_info(workbook) -> dict[str, str]:
    sheet = workbook["AllocationInfo"]
    headers = _require_headers(sheet, {"Field", "Value"})
    result = {}
    for row in range(2, sheet.max_row + 1):
        field = sheet.cell(row, headers["Field"]).value
        value = sheet.cell(row, headers["Value"]).value
        if field is not None:
            result[str(field).strip()] = value
    return result

def _parse_revision_workbook(file_path: str | Path, expected_parent_id: str) -> tuple[list[dict], list[dict]]:
    workbook = load_workbook(file_path, data_only=True)
    missing_sheets = sorted(REQUIRED_REVISION_SHEETS - set(workbook.sheetnames))
    if missing_sheets:
        raise ValueError(f"Revision workbook is missing required sheet(s): {', '.join(missing_sheets)}.")
    info = _read_allocation_info(workbook)
    workbook_parent = _normalize_id(info.get("ParentAllocationID"))
    if workbook_parent != _normalize_id(expected_parent_id):
        raise ValueError(f"Workbook ParentAllocationID '{workbook_parent}' does not match requested parent allocation '{expected_parent_id}'.")
    team_sheet = workbook["TeamAssignments"]
    team_headers = _require_headers(team_sheet, {"StudentID", "TeamNumber", "ProjectID"})
    team_rows = []
    for row in range(2, team_sheet.max_row + 1):
        student_id = _normalize_id(team_sheet.cell(row, team_headers["StudentID"]).value)
        project_id = _normalize_id(team_sheet.cell(row, team_headers["ProjectID"]).value)
        team_value = team_sheet.cell(row, team_headers["TeamNumber"]).value
        if not student_id and not project_id and team_value in (None, ""):
            continue
        if not student_id or not project_id:
            raise ValueError(f"TeamAssignments row {row} must contain StudentID, TeamNumber and ProjectID.")
        team_rows.append({"student_id": student_id, "team_number": _as_int(team_value, f"TeamAssignments row {row} TeamNumber"), "project_id": project_id})
    supervisor_sheet = workbook["SupervisorAssignments"]
    supervisor_headers = _require_headers(supervisor_sheet, {"TeamNumber", "ProjectID", "SupervisorID"})
    supervisor_rows = []
    for row in range(2, supervisor_sheet.max_row + 1):
        project_id = _normalize_id(supervisor_sheet.cell(row, supervisor_headers["ProjectID"]).value)
        supervisor_id = _normalize_id(supervisor_sheet.cell(row, supervisor_headers["SupervisorID"]).value)
        team_value = supervisor_sheet.cell(row, supervisor_headers["TeamNumber"]).value
        if not project_id and not supervisor_id and team_value in (None, ""):
            continue
        if not project_id or not supervisor_id:
            raise ValueError(f"SupervisorAssignments row {row} must contain TeamNumber, ProjectID and SupervisorID.")
        supervisor_rows.append({"team_number": _as_int(team_value, f"SupervisorAssignments row {row} TeamNumber"), "project_id": project_id, "supervisor_id": supervisor_id})
    return team_rows, supervisor_rows

def _preference_dissatisfaction(ranked_projects: list[str], project_id: str, project_count: int) -> tuple[int | None, float]:
    ranked = [_normalize_id(item) for item in ranked_projects if _normalize_id(item)]
    if project_count <= 1:
        return (1 if project_id in ranked else None), 0.0
    if project_id not in ranked:
        return None, 1.0
    rank = ranked.index(project_id) + 1
    if len(ranked) >= project_count:
        denominator = project_count - 1
    else:
        denominator = max(len(ranked), 1)
    return rank, (rank - 1) / denominator

def _supervisor_assignment_details(reference_data: dict, project_id: str, supervisor_id: str, assignment_counts: Counter) -> dict:
    supervisors = {_normalize_id(item.get("supervisor_id")): item for item in reference_data.get("supervisors") or []}
    if supervisor_id not in supervisors:
        raise ValueError(f"Unknown supervisor '{supervisor_id}'.")
    supervisor = supervisors[supervisor_id]
    current_load = int(supervisor.get("current_load") or 0)
    maximum_teams = int(supervisor.get("maximum_teams") or 0)
    assigned_count = assignment_counts[supervisor_id]
    projected_load = current_load + assigned_count
    if projected_load > maximum_teams:
        raise ValueError(f"Supervisor {supervisor_id} exceeds capacity: current load {current_load} + revised assignments {assigned_count} > maximum {maximum_teams}.")
    project_domains = sorted({str(item.get("domain") or "").strip() for item in reference_data.get("project_domains") or [] if _normalize_id(item.get("project_id")) == project_id and str(item.get("domain") or "").strip()})
    expertise = {str(item.get("domain") or "").strip().casefold() for item in reference_data.get("supervisor_domains") or [] if _normalize_id(item.get("supervisor_id")) == supervisor_id and str(item.get("type") or "").strip().casefold() == "expertise" and str(item.get("domain") or "").strip()}
    interests = {str(item.get("domain") or "").strip().casefold() for item in reference_data.get("supervisor_domains") or [] if _normalize_id(item.get("supervisor_id")) == supervisor_id and str(item.get("type") or "").strip().casefold() == "interest" and str(item.get("domain") or "").strip()}
    expertise_matches = [domain for domain in project_domains if domain.casefold() in expertise]
    interest_matches = [domain for domain in project_domains if domain.casefold() in interests]
    if project_domains and all(domain.casefold() in expertise for domain in project_domains):
        match_status = "Strong Expertise Match"
    elif expertise_matches:
        match_status = "Expertise Match"
    elif interest_matches:
        match_status = "Interest Match"
    else:
        match_status = "Capacity-Only Assignment"
    explanation_parts = []
    if expertise_matches:
        explanation_parts.append(f"Expertise match: {', '.join(expertise_matches)}.")
    if interest_matches:
        explanation_parts.append(f"Interest match: {', '.join(interest_matches)}.")
    if not explanation_parts:
        explanation_parts.append("No project-domain expertise or interest match was identified.")
    explanation_parts.append("Capacity constraint satisfied.")
    return {
        "project_id": project_id,
        "supervisor_id": supervisor_id,
        "supervisor_name": str(supervisor.get("supervisor_name") or ""),
        "match_status": match_status,
        "current_load": current_load,
        "maximum_teams": maximum_teams,
        "projected_load": projected_load,
        "remaining_capacity": maximum_teams - projected_load,
        "project_domains": project_domains,
        "expertise_matches": expertise_matches,
        "interest_matches": interest_matches,
        "explanation": " ".join(explanation_parts),
    }

def _build_revision_candidate(parent: dict, reference_data: dict, team_rows: list[dict], supervisor_rows: list[dict]) -> tuple[dict, list[dict], list[str]]:
    students = {_normalize_id(item.get("student_id")): item for item in reference_data.get("students") or []}
    preferences = {_normalize_id(item.get("student_id")): [_normalize_id(p) for p in item.get("ranked_projects") or []] for item in reference_data.get("preferences") or []}
    projects = {_normalize_id(item.get("project_id")): item for item in reference_data.get("projects") or []}
    requirements_by_project = defaultdict(list)
    for requirement in reference_data.get("project_requirements") or []:
        requirements_by_project[_normalize_id(requirement.get("project_id"))].append(requirement)
    provided_students = [row["student_id"] for row in team_rows]
    duplicates = sorted(student_id for student_id, count in Counter(provided_students).items() if count > 1)
    if duplicates:
        raise ValueError(f"Duplicate student assignment(s): {', '.join(duplicates[:10])}.")
    expected_students = set(students)
    provided_student_set = set(provided_students)
    if provided_student_set != expected_students:
        missing = sorted(expected_students - provided_student_set)
        extra = sorted(provided_student_set - expected_students)
        raise ValueError(f"Revised allocation must contain every original student exactly once. Missing={missing[:10]}, unexpected={extra[:10]}.")
    grouped = defaultdict(list)
    team_project = {}
    for row in team_rows:
        if row["team_number"] <= 0:
            raise ValueError("TeamNumber values must be positive integers.")
        if row["project_id"] not in projects:
            raise ValueError(f"Unknown project '{row['project_id']}'.")
        grouped[row["team_number"]].append(row["student_id"])
        existing_project = team_project.setdefault(row["team_number"], row["project_id"])
        if existing_project != row["project_id"]:
            raise ValueError(f"Team {row['team_number']} contains more than one ProjectID.")
    student_count = len(students)
    team_size = int(parent["students_per_team"])
    required_team_count = ceil(student_count / team_size)
    expected_team_numbers = set(range(1, required_team_count + 1))
    if set(grouped) != expected_team_numbers:
        raise ValueError(f"Team numbers must be exactly 1 to {required_team_count}.")
    project_ids = list(team_project.values())
    if len(project_ids) != len(set(project_ids)):
        raise ValueError("Each project may be assigned to only one team.")
    if set(project_ids) != set(projects):
        missing = sorted(set(projects) - set(project_ids))
        extra = sorted(set(project_ids) - set(projects))
        raise ValueError(f"Revised allocation must use exactly the approved project set. Missing={missing}, unexpected={extra}.")
    quotient, remainder = divmod(student_count, team_size)
    expected_sizes = sorted(([team_size] * quotient) + ([remainder] if remainder else []))
    actual_sizes = sorted(len(members) for members in grouped.values())
    if actual_sizes != expected_sizes:
        raise ValueError(f"Invalid revised team sizes. Expected sorted team sizes {expected_sizes}, received {actual_sizes}.")
    warnings = []
    if remainder == 1:
        warnings.append("The revised allocation contains the permitted one-student remainder team.")
    supervisor_by_team = {}
    for row in supervisor_rows:
        if row["team_number"] in supervisor_by_team:
            raise ValueError(f"SupervisorAssignments contains duplicate TeamNumber {row['team_number']}.")
        supervisor_by_team[row["team_number"]] = row
    if set(supervisor_by_team) != expected_team_numbers:
        raise ValueError("SupervisorAssignments must contain exactly one row for every team.")
    for team_number, row in supervisor_by_team.items():
        if row["project_id"] != team_project[team_number]:
            raise ValueError(f"SupervisorAssignments team {team_number} ProjectID must match TeamAssignments.")
    assignment_counts = Counter(row["supervisor_id"] for row in supervisor_rows)
    teams = []
    all_dissatisfaction = []
    supervisor_assignments = []
    for team_number in sorted(grouped):
        project_id = team_project[team_number]
        project = projects[project_id]
        member_ids = sorted(grouped[team_number])
        project_requirements = requirements_by_project[project_id]
        requirement_results = []
        relevant_technologies = [str(item.get("technology") or "").strip() for item in project_requirements]
        for requirement in project_requirements:
            technology = str(requirement.get("technology") or "").strip()
            min_level = int(requirement.get("min_level") or 0)
            required_members = int(requirement.get("required_members") or 0)
            qualified = [student_id for student_id in member_ids if int((students[student_id].get("skills") or {}).get(technology, 0)) >= min_level]
            coverage = 1.0 if required_members <= 0 else min(1.0, len(qualified) / required_members)
            requirement_results.append({
                "technology": technology,
                "min_level": min_level,
                "required_members": required_members,
                "qualified_members": len(qualified),
                "qualified_students": qualified,
                "coverage": coverage,
                "status": "Covered" if coverage >= 1.0 else "Gap",
            })
        technical_coverage = sum(item["coverage"] for item in requirement_results) / len(requirement_results) if requirement_results else 1.0
        technical_deficit = 1.0 - technical_coverage
        student_results = []
        first_choice_count = 0
        ranked_choice_count = 0
        unranked_count = 0
        team_dissatisfaction = []
        for student_id in member_ids:
            ranked_projects = preferences.get(student_id, [])
            preference_rank, dissatisfaction = _preference_dissatisfaction(ranked_projects, project_id, len(projects))
            if preference_rank == 1:
                first_choice_count += 1
            if preference_rank is None:
                unranked_count += 1
            else:
                ranked_choice_count += 1
            team_dissatisfaction.append(dissatisfaction)
            all_dissatisfaction.append(dissatisfaction)
            skills = students[student_id].get("skills") or {}
            student_results.append({
                "student_id": student_id,
                "preference_rank": preference_rank,
                "dissatisfaction": dissatisfaction,
                "ranked_projects": ranked_projects,
                "relevant_skills": {technology: int(skills.get(technology, 0)) for technology in relevant_technologies},
            })
        average_dissatisfaction = sum(team_dissatisfaction) / len(team_dissatisfaction) if team_dissatisfaction else 0.0
        teams.append({
            "team_number": team_number,
            "project_id": project_id,
            "project_title": str(project.get("project_title") or ""),
            "team_size": len(member_ids),
            "technical_coverage": technical_coverage,
            "technical_deficit": technical_deficit,
            "technical_status": "Fully covered" if technical_deficit <= 1e-12 else "Partially covered",
            "preference_summary": {
                "average_dissatisfaction": average_dissatisfaction,
                "first_choice_count": first_choice_count,
                "ranked_choice_count": ranked_choice_count,
                "unranked_count": unranked_count,
            },
            "students": student_results,
            "requirements": requirement_results,
        })
        supervisor_row = supervisor_by_team[team_number]
        supervisor_assignments.append(_supervisor_assignment_details(reference_data, project_id, supervisor_row["supervisor_id"], assignment_counts))
    global_deficit = sum(team["technical_deficit"] for team in teams) / len(teams) if teams else 0.0
    global_dissatisfaction = sum(all_dissatisfaction) / len(all_dissatisfaction) if all_dissatisfaction else 0.0
    solution = {
        "solution_id": int(parent.get("solution_id") or 0),
        "role": f"Staff-revised allocation (Revision {int(parent.get('revision_number') or 1) + 1})",
        "technical_requirement_deficit": global_deficit,
        "technical_requirement_coverage": 1.0 - global_deficit,
        "preference_dissatisfaction": global_dissatisfaction,
        "preference_satisfaction": 1.0 - global_dissatisfaction,
        "integrity": {"valid": True},
        "teams": teams,
    }
    return solution, supervisor_assignments, warnings

def _change_summary(parent: dict, solution: dict, supervisor_assignments: list[dict]) -> dict:
    old_student_map = {}
    for team in parent.get("teams") or []:
        for student in team.get("students") or []:
            old_student_map[student["student_id"]] = {"team_number": team["team_number"], "project_id": team["project_id"]}
    new_student_map = {}
    for team in solution.get("teams") or []:
        for student in team.get("students") or []:
            new_student_map[student["student_id"]] = {"team_number": team["team_number"], "project_id": team["project_id"]}
    moved_students = []
    for student_id in sorted(old_student_map):
        if old_student_map[student_id] != new_student_map.get(student_id):
            moved_students.append({"student_id": student_id, "from": old_student_map[student_id], "to": new_student_map.get(student_id)})
    old_supervisors = {team["project_id"]: (team.get("supervisor") or {}).get("supervisor_id") for team in parent.get("teams") or []}
    new_supervisors = {item["project_id"]: item["supervisor_id"] for item in supervisor_assignments}
    changed_supervisors = []
    for project_id in sorted(old_supervisors):
        if old_supervisors[project_id] != new_supervisors.get(project_id):
            changed_supervisors.append({"project_id": project_id, "from_supervisor_id": old_supervisors[project_id], "to_supervisor_id": new_supervisors.get(project_id)})
    return {
        "moved_student_count": len(moved_students),
        "moved_students": moved_students,
        "changed_supervisor_count": len(changed_supervisors),
        "changed_supervisors": changed_supervisors,
        "original_metrics": {
            "technical_requirement_deficit": parent["technical_requirement_deficit"],
            "technical_requirement_coverage": parent["technical_requirement_coverage"],
            "preference_dissatisfaction": parent["preference_dissatisfaction"],
            "preference_satisfaction": parent["preference_satisfaction"],
        },
        "revised_metrics": {
            "technical_requirement_deficit": solution["technical_requirement_deficit"],
            "technical_requirement_coverage": solution["technical_requirement_coverage"],
            "preference_dissatisfaction": solution["preference_dissatisfaction"],
            "preference_satisfaction": solution["preference_satisfaction"],
        },
    }

def validate_revision_workbook(db: Session, parent_allocation_id: str, file_path: str | Path) -> dict:
    parent_model = db.scalar(select(FinalAllocation).where(FinalAllocation.id == parent_allocation_id))
    if parent_model is None:
        raise ValueError("Parent final allocation was not found.")
    if parent_model.status != "ACTIVE":
        raise ValueError("Only the current ACTIVE allocation can be revised. Download a fresh editable workbook from the current ACTIVE allocation.")
    if not parent_model.reference_data:
        raise ValueError("The parent allocation does not contain source reference data required for safe recalculation.")
    parent = get_final_allocation(db, parent_allocation_id)
    team_rows, supervisor_rows = _parse_revision_workbook(file_path, parent_allocation_id)
    solution, supervisor_assignments, warnings = _build_revision_candidate(parent, parent_model.reference_data, team_rows, supervisor_rows)
    changes = _change_summary(parent, solution, supervisor_assignments)
    if changes["moved_student_count"] == 0 and changes["changed_supervisor_count"] == 0:
        raise ValueError("No allocation changes were detected. Edit team/project assignments or supervisor assignments before creating a revision.")
    return {
        "valid": True,
        "parent_allocation_id": parent_allocation_id,
        "next_revision_number": int(parent_model.revision_number or 1) + 1,
        "warnings": warnings,
        "changes": changes,
        "candidate": {
            "student_count": parent["student_count"],
            "team_count": len(solution["teams"]),
            "technical_requirement_deficit": solution["technical_requirement_deficit"],
            "technical_requirement_coverage": solution["technical_requirement_coverage"],
            "preference_dissatisfaction": solution["preference_dissatisfaction"],
            "preference_satisfaction": solution["preference_satisfaction"],
        },
        "_solution": solution,
        "_supervisor_assignments": supervisor_assignments,
    }

def create_revised_allocation(db: Session, parent_allocation_id: str, file_path: str | Path, source_file_name: str | None, change_reason: str | None) -> dict:
    preview = validate_revision_workbook(db, parent_allocation_id, file_path)
    parent_model = db.scalar(select(FinalAllocation).where(FinalAllocation.id == parent_allocation_id))
    if parent_model is None or parent_model.status != "ACTIVE":
        raise ValueError("The parent allocation is no longer ACTIVE. Download a new editable workbook and validate again.")
    parent = get_final_allocation(db, parent_allocation_id)
    reason = (change_reason or "").strip() or "Manual staff revision uploaded through the allocation revision workflow."
    revised = persist_final_allocation(
        db,
        source_file_name=source_file_name,
        students_per_team=parent["students_per_team"],
        algorithm=parent["algorithm"],
        optimizer_version=parent["optimizer_version"],
        solution=preview.pop("_solution"),
        supervisor_assignments=preview.pop("_supervisor_assignments"),
        reference_data=parent_model.reference_data,
        allocation_source="MANUAL_REVISION",
        revision_number=int(parent_model.revision_number or 1) + 1,
        parent_allocation_id=parent_allocation_id,
        change_reason=reason,
    )
    return {"allocation": revised, "changes": preview["changes"], "warnings": preview["warnings"]}
