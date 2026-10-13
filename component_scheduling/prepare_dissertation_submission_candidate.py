"""Reconcile the reviewed dissertation with the current, retained evidence.

This deliberately reads the earlier submission-review document, not the later
Research_Corrected copy, which reintroduced obsolete route/capacity claims and
an acknowledgement copied from a different component. Both inputs are kept.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "IT22117014_Individual_Component_Dissertation_Submission_Review.docx"
OUTPUT = ROOT / "IT22117014_Individual_Component_Dissertation_Submission_Candidate.docx"
TITLE = "Adaptive Scheduling and Generative Project Planning"


def replace_paragraph(doc: Document, prefix: str, replacement: str) -> None:
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}; found {len(matches)}")
    matches[0].text = replacement


def set_cell_text(table, row: int, column: int, value: str) -> None:
    """Keep the dataset table's existing compact type size."""
    paragraph = table.cell(row, column).paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run(value)
    run.font.size = Pt(8.5)


def make() -> None:
    doc = Document(SOURCE)

    replace_paragraph(doc, "Generative Project Planning and Adaptive Scheduling with", TITLE)
    replace_paragraph(
        doc,
        "Undergraduate project teams must turn proposals into task plans, allocate work, inspect evidence and arrange academic evaluation events. This dissertation examines",
        "Undergraduate project teams must turn proposals into task plans, allocate work, inspect evidence and arrange academic evaluation events. This dissertation examines the Adaptive Scheduling and Generative Project Planning component of the Intelligent Project Management System (IPMS). Its parser, effort estimator, dependency-aware Gantt generator, evidence records and academic timetable form an evidence-aware prototype. The live workload transition uses Longest Processing Time (LPT) assignment and persists a separate balanced schedule version. An isolated-PostgreSQL regression run on 11 October 2026 passed 11 focused checks, including hall-capacity enforcement, schedule-history member restoration and a synthetic local PDF approval and signed-webhook path. In 100 simulated historical-profile trials, LPT averaged 3.72% planned-workload CV, while a historical classifier-guided hybrid averaged 4.40%. A legacy mixed-label JIRA row split reported 81.27% accuracy, but a chronological GFG-only proxy test achieved 31.94%, below its 77.31% majority baseline; neither result measures student suitability. Rule extraction recovered 73/75 listed tasks in synthetic proposals, not general extraction accuracy. The effort regressor's log MAE was 0.8554 and its measured original-hour MAPE about 152.6%. The contribution is a tested local planning workflow with explicit evidence boundaries, not demonstrated improvements in real-cohort outcomes.",
    )

    replace_paragraph(
        doc,
        "The SciREX and world-cities files in the workspace are not training inputs",
        "The workload experiment reads two local files. Despite its legacy filename, GFG_FINAL.csv contains 49,000 Sourcetree-for-Windows (SRCTREEWIN) Jira rows; filtering for an assignee and parseable creation date leaves 1,078 issues from five assignee IDs. A matching issue key is visible in Atlassian's public tracker [17], but the exact bulk-export query, acquisition date and reuse terms are not recorded. The second file, apache_jira_delay.csv, contains 333 prepared rows bearing five Apache project labels. It lacks issue keys, source URLs and an extraction manifest, so its row-level origin and transformation cannot be reconstructed; project codes serve as pseudo-members, not people. The effort file combines 499 China records and 81 Desharnais records. The China adjusted-function-point/effort pairs match all 499 records, in order, in Zenodo's China v1 release [18]. The Desharnais pairs match the local 81-row Desharnais file and a publicly available PROMISE-labelled copy, but the exact acquisition route is undocumented [16]. These are software-project proxies, not student outcomes. SciREX, world-cities, panel sheets and prepared commit/document fixtures are not inputs to these evaluated models. Redistribution rights still require a separate provider check.",
    )

    # Replace the misleading legacy dataset nickname in explanatory prose, but
    # preserve the actual on-disk filename for reproducibility.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            original = run.text
            updated = original.replace("GFG-only", "Sourcetree-only")
            updated = updated.replace("GFG ", "Sourcetree ")
            updated = updated.replace("GFG.", "Sourcetree.")
            if updated != original:
                run.text = updated

    dataset_table = next(
        table for table in doc.tables
        if table.cell(0, 0).text == "Data source" and table.cell(0, 1).text == "Size"
    )
    set_cell_text(dataset_table, 1, 0, "Sourcetree Jira CSV (GFG_FINAL.csv)")
    set_cell_text(dataset_table, 1, 1, "49,000 raw; 1,078 used")
    set_cell_text(dataset_table, 1, 2, "Assigned, dated issue-status proxy; five assignees")
    set_cell_text(dataset_table, 2, 0, "Local Apache-project summary")
    set_cell_text(dataset_table, 2, 1, "333 prepared rows")
    set_cell_text(dataset_table, 2, 2, "Issue-derived proxy; origin not row-verifiable")
    set_cell_text(dataset_table, 3, 0, "China + Desharnais local merge")
    set_cell_text(dataset_table, 3, 1, "499 + 81 projects")
    set_cell_text(dataset_table, 3, 2, "Effort regression; China pairs checked vs. Zenodo [18]")
    for table in doc.tables:
        if table is dataset_table:
            continue
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        original = run.text
                        updated = original.replace("GFG-only", "Sourcetree-only")
                        updated = updated.replace("GFG ", "Sourcetree ")
                        if updated != original:
                            run.text = updated
    replace_paragraph(
        doc,
        "An older audit mentions 60/60 backend tests",
        "An older audit mentions 60/60 backend tests, but its raw run log was not retained and is not counted as current evidence. A focused run on 11 October 2026 passed 11 tests against the isolated PostgreSQL database ipms_scheduling_codex_test_20261010. It covered the unique authenticated LPT route, versioning and rollback, production demo-auth refusal, hall-seat rejection, schedule-history member IDs, mixed-format deadline scans and one synthetic local PDF submission → supervisor approval → signed task-key webhook journey. The fixture report separately verifies 30 panel records, eight hall records, 13 prepared task-keyed commits and five example documents. These checks do not establish real student authentication, a new GitHub.com webhook delivery, cross-component payload integration or improved academic outcomes.",
    )
    replace_paragraph(
        doc,
        "The older 60-test claim has no retained raw output",
        "The older 60-test claim has no retained raw output and is not counted as current evidence. The 11 focused PostgreSQL checks establish selected assignment, versioning, hall-capacity, History, deadline and synthetic local document/webhook behaviours. They do not test every negative review decision, all UI states, direct Team Formation or Quality/Risk payload exchange, or real-cohort outcomes. PDF fixtures are not scanned or varied real proposals; panel sheets are synthetic; and the GitHub test repository uses known task keys. These limits accompany each reported result.",
    )
    replace_paragraph(
        doc,
        "Review tests should check both positive and negative decisions.",
        "Review tests should check both positive and negative decisions. The focused local run covered one approved PDF submission and a signed, task-keyed webhook, with persisted review and task-progress records. Rejection without a comment, unauthorised reviewer access, malformed files, duplicate delivery and unmatched-commit paths remain separate acceptance cases unless their execution logs are retained. The local signed payload is not a fresh GitHub.com delivery, and these checks do not amount to a production student-login test.",
    )
    replace_paragraph(
        doc,
        "Numerical results in this dissertation come from saved training or offline-evaluation outputs",
        "Numerical results in this dissertation come from saved training or offline-evaluation outputs, a stated recomputation on the same hold-out, and the retained 11 October 2026 focused database-test record. The WBS count and sample workbooks remain controlled fixtures. The 11 checks cover selected live scheduling, History, deadline, hall-capacity and synthetic local PDF-review/webhook paths; they do not validate every API, cross-component deployment or real student outcome. Proposed negative tests outside this focused suite remain acceptance criteria. Any later code or dataset change requires a fresh reproducible run before submission.",
    )

    # The final appendicial evidence map is a two-column table. Add the latest
    # retained test/audit records without implying that they prove a full rollout.
    evidence_table = doc.tables[-1]
    for label, location in [
        ("11 Oct. regression", "component_scheduling/Scheduling_Submission_Verification_2026-10-11.md"),
        ("Synthetic-flow detail", "component_scheduling/Scheduling_Fresh_Integration_Test_Evidence_2026-10-10.md"),
        ("External commit audit", "component_scheduling/External_Evidence_Audit_2026-10-10.md"),
        ("Dataset provenance", "component_scheduling/Dataset_Source_Provenance_Audit_2026-10-11.md"),
    ]:
        cells = evidence_table.add_row().cells
        cells[0].text = label
        cells[1].text = location
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    appendix = next(p for p in doc.paragraphs if p.text == "Appendix A. Evidence and Code Map")
    appendix.insert_paragraph_before(
        '[17] Atlassian, "SRCTREEWIN-14001: Authentication failed when attempting Fetch command," Sourcetree for Windows Jira issue, 2023. [Online]. Available: https://jira.atlassian.com/browse/SRCTREEWIN-14001',
        style="Normal",
    )
    appendix.insert_paragraph_before(
        '[18] F. H. Yun, "China: Effort Estimation Dataset," Zenodo, v1, 2010. doi: 10.5281/zenodo.268446.',
        style="Normal",
    )

    for section in doc.sections:
        if section.header.paragraphs and "Generative Planning and Adaptive Scheduling" in section.header.paragraphs[0].text:
            section.header.paragraphs[0].text = "IPMS  |  Adaptive Scheduling and Generative Project Planning"
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Individual component dissertation; author and institutional review pending"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    make()
