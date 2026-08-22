"""
Generate Excel templates for the Academic Event Auto-Scheduler.
Run this script once to create the template files.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── Supervisors Template ────────────────────────────────────────────────────
wb_sup = openpyxl.Workbook()
ws = wb_sup.active
ws.title = "Supervisors"

headers = ["Name", "Email", "Research Expertise (keywords)", "Available Dates (comma separated)", "Available Times (HH:MM-HH:MM)"]
sample_rows = [
    ["Dr. A. Perera", "a.perera@sliit.lk", "machine learning, artificial intelligence, NLP, data science, deep learning", "2026-10-01,2026-10-02,2026-10-03", "09:00-12:00"],
    ["Dr. B. Silva",  "b.silva@sliit.lk",  "software engineering, agile methodology, project management, scrum, devops", "2026-10-01,2026-10-04,2026-10-05", "13:00-17:00"],
    ["Dr. C. Fernando","c.fernando@sliit.lk","computer vision, image processing, deep learning, convolutional neural networks", "2026-10-02,2026-10-03,2026-10-06", "09:00-13:00"],
    ["Dr. D. Jayawardena","d.jayawardena@sliit.lk","cybersecurity, network security, cryptography, blockchain", "2026-10-04,2026-10-05,2026-10-07", "10:00-15:00"],
    ["Dr. E. Wijesinghe","e.wijesinghe@sliit.lk","database systems, cloud computing, distributed systems, big data", "2026-10-01,2026-10-06,2026-10-08", "09:00-12:00"],
]

header_fill = PatternFill("solid", fgColor="1E3A5F")
header_font = Font(bold=True, color="FFFFFF", size=11)
border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'),  bottom=Side(style='thin')
)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill  = header_fill
    cell.font  = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

row_fills = ["EEF2FF", "F8FAFC"]
for r, row in enumerate(sample_rows, 2):
    fill = PatternFill("solid", fgColor=row_fills[r % 2])
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.fill      = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border    = border

col_widths = [22, 28, 55, 38, 25]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30

wb_sup.save("supervisors_template.xlsx")
print("OK: supervisors_template.xlsx created")


# ─── Halls Template ───────────────────────────────────────────────────────────
wb_hall = openpyxl.Workbook()
ws2 = wb_hall.active
ws2.title = "Halls"

headers2 = ["Hall Name", "Capacity", "Floor / Building", "Available Dates (comma separated)", "Available Times (HH:MM-HH:MM)"]
sample_rows2 = [
    ["Lab 301",         30,  "3rd Floor, Block A", "2026-10-01,2026-10-02,2026-10-03,2026-10-04", "09:00-17:00"],
    ["Lab 302",         30,  "3rd Floor, Block A", "2026-10-01,2026-10-05,2026-10-06",             "09:00-17:00"],
    ["Lecture Hall A",  60,  "1st Floor, Block B", "2026-10-02,2026-10-03,2026-10-07",             "08:00-12:00"],
    ["Lecture Hall B",  80,  "2nd Floor, Block B", "2026-10-04,2026-10-05,2026-10-08",             "13:00-17:00"],
    ["Auditorium",      200, "Ground Floor, Main", "2026-10-01,2026-10-06,2026-10-09",             "09:00-13:00"],
    ["Seminar Room 1",  20,  "4th Floor, Block C", "2026-10-02,2026-10-04,2026-10-07",             "10:00-16:00"],
    ["Online - Zoom",   999, "Virtual",            "2026-10-01,2026-10-02,2026-10-03,2026-10-04,2026-10-05,2026-10-06,2026-10-07,2026-10-08,2026-10-09", "08:00-20:00"],
]

for col, h in enumerate(headers2, 1):
    cell = ws2.cell(row=1, column=col, value=h)
    cell.fill      = PatternFill("solid", fgColor="14532D")
    cell.font      = Font(bold=True, color="FFFFFF", size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border

for r, row in enumerate(sample_rows2, 2):
    fill = PatternFill("solid", fgColor=["F0FDF4","DCFCE7"][r % 2])
    for c, val in enumerate(row, 1):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.fill      = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border    = border

col_widths2 = [20, 12, 25, 50, 25]
for i, w in enumerate(col_widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.row_dimensions[1].height = 30

wb_hall.save("halls_template.xlsx")
print("OK: halls_template.xlsx created")
print("Fill in your real data in these files, then upload them in the Academic Scheduler tab.")
