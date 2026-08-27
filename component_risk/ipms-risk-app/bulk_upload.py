"""Bulk team creation: a supervisor uploads one Excel file (one row per member)
and the system creates the projects, the student accounts and the team links."""
import io
import re

from openpyxl import Workbook, load_workbook

import auth
import database

# Header names are matched after removing everything except letters, so
# "IT Number", "it_number" and "IT Numbers" all work.
HEADER_ALIASES = {
    "team_id": {"projectid", "teamid"},
    "name": {"projectname"},
    "leader": {"teamleader", "leader"},
    "member_name": {"membername", "membersname", "membersnames", "member", "studentname"},
    "it_number": {"itnumber", "itnumbers", "itno"},
    "email": {"itemail", "itemails", "itmail", "itmails", "email", "studentemail"},
    "github_login": {"githubusername", "githubuser", "githublogin"},
    "github_url": {"githubrepourl", "githuburl", "githubrepo", "repourl", "repositoryurl"},
    "jira": {"jira", "jiraprojectkey", "jirakey", "jiraoptional"}
}

REQUIRED_COLUMNS = ("team_id", "name", "member_name", "email")
LEADER_VALUES = {"yes", "y", "true", "1", "x", "leader"}

TEMPLATE_HEADERS = [
    "Project ID", "Project Name", "Team Leader", "Member Name", "IT Number",
    "IT Email", "GitHub Username", "GitHub Repo URL", "Jira Project Key"
]
TEMPLATE_ROWS = [
    ["TEAM-101", "Smart Campus App", "Yes", "Kasun Silva", "IT22109576",
     "it22109576@my.sliit.lk", "kasun-dev", "https://github.com/team101/smart-campus", "SCA"],
    ["TEAM-101", "Smart Campus App", "", "Amaya Fernando", "IT22110001",
     "it22110001@my.sliit.lk", "amaya-f", "https://github.com/team101/smart-campus", "SCA"],
    ["TEAM-102", "Library Portal", "Yes", "Nuwan Perera", "IT22110234",
     "it22110234@my.sliit.lk", "", "https://github.com/team102/library-portal", ""]
]


def build_template():
    """Returns the downloadable .xlsx template as bytes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Teams"

    sheet.append(TEMPLATE_HEADERS)
    for row in TEMPLATE_ROWS:
        sheet.append(row)

    widths = [12, 22, 12, 20, 14, 26, 18, 42, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def normalize_header(value):
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def map_headers(header_row):
    """Maps column index -> field name for every recognised header."""
    mapping = {}
    for index, cell in enumerate(header_row):
        compact = normalize_header(cell)
        for field, aliases in HEADER_ALIASES.items():
            if compact in aliases:
                mapping[index] = field
                break
    return mapping


def cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_workbook(content):
    """Reads the uploaded bytes into project groups.

    Returns (groups, errors): groups is a list of dicts with project fields
    and a members list; errors is a list of row-level problem descriptions.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        return [], ["Could not read the file — make sure it is a .xlsx Excel file."]

    rows = list(workbook.active.iter_rows(values_only=True))

    header_index = None
    columns = {}
    for index, row in enumerate(rows):
        candidate = map_headers(row)
        if len(candidate) >= 3:
            header_index, columns = index, candidate
            break

    if header_index is None:
        return [], ["No header row found. Download the template to see the expected columns."]

    mapped_fields = set(columns.values())
    missing = [field for field in REQUIRED_COLUMNS if field not in mapped_fields]
    if missing:
        labels = {"team_id": "Project ID", "name": "Project Name",
                  "member_name": "Member Name", "email": "IT Email"}
        return [], ["Missing required column(s): " + ", ".join(labels[f] for f in missing)]

    groups = {}
    errors = []

    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        values = {field: "" for field in HEADER_ALIASES}
        for index, field in columns.items():
            if index < len(row):
                values[field] = cell_text(row[index])

        if not any(values.values()):
            continue  # fully empty row

        if not values["team_id"]:
            errors.append(f"Row {row_number}: Project ID is missing — row skipped.")
            continue

        group = groups.setdefault(values["team_id"], {
            "team_id": values["team_id"],
            "name": "",
            "github_url": "",
            "jira": "",
            "members": [],
            "warnings": []
        })

        # Project-level fields: first non-empty value wins.
        for field in ("name", "github_url", "jira"):
            if values[field] and not group[field]:
                group[field] = values[field]

        if not values["member_name"]:
            errors.append(f"Row {row_number}: Member Name is missing — row skipped.")
            continue

        if "@" not in values["email"]:
            errors.append(
                f"Row {row_number} ({values['member_name']}): "
                "a valid IT Email is required — row skipped."
            )
            continue

        group["members"].append({
            "row": row_number,
            "name": values["member_name"],
            "email": values["email"].lower(),
            "it_number": values["it_number"],
            "github_login": values["github_login"],
            "leader": values["leader"].lower() in LEADER_VALUES
        })

    return list(groups.values()), errors


def apply_upload(groups, supervisor_id):
    """Creates/updates projects, accounts and memberships. Returns report entries."""
    entries = []

    for group in groups:
        entry = {
            "team_id": group["team_id"],
            "name": group["name"] or group["team_id"],
            "status": "",
            "members_added": 0,
            "accounts_created": [],
            "problems": list(group["warnings"])
        }

        existing = database.get_project_by_team_id(group["team_id"])

        if existing and existing["supervisor_id"] != supervisor_id:
            entry["status"] = "skipped — this Project ID belongs to another supervisor"
            entries.append(entry)
            continue

        if existing:
            database.update_project(
                existing["id"],
                group["name"] or existing["name"],
                group["team_id"],
                group["github_url"] or existing["github_url"],
                group["jira"] or existing["jira_project_key"]
            )
            project_id = existing["id"]
            entry["status"] = "updated"
        else:
            if not group["name"]:
                entry["status"] = "skipped — Project Name is missing"
                entries.append(entry)
                continue
            project_id = database.create_project(
                group["name"], group["team_id"],
                group["github_url"], group["jira"], supervisor_id
            )
            entry["status"] = "created"

        if not group["github_url"] and not (existing and existing["github_url"]):
            entry["problems"].append(
                "No GitHub Repo URL — analytics and risk prediction need one."
            )

        leader_assigned = False
        for member in group["members"]:
            user = database.get_user_by_email(member["email"])

            if user is None:
                if not member["it_number"]:
                    entry["problems"].append(
                        f"{member['name']} (row {member['row']}): new accounts need an "
                        "IT Number (it becomes their first password) — member skipped."
                    )
                    continue
                password_hash, salt = auth.hash_password(member["it_number"])
                user_id = database.create_user(
                    member["email"], member["name"], password_hash, salt,
                    "student", member["it_number"]
                )
                entry["accounts_created"].append(member["email"])
            else:
                if user["role"] != "student":
                    entry["problems"].append(
                        f"{member['name']} (row {member['row']}): {member['email']} is a "
                        "supervisor account — member skipped."
                    )
                    continue
                user_id = user["id"]
                if member["it_number"] and not user["it_number"]:
                    database.set_user_it_number(user_id, member["it_number"])

            is_leader = 0
            if member["leader"]:
                if leader_assigned:
                    entry["problems"].append(
                        f"Multiple team leaders marked — kept the first, "
                        f"{member['name']} stays a regular member."
                    )
                else:
                    is_leader = 1
                    leader_assigned = True

            database.add_member(
                project_id, user_id,
                github_login=member["github_login"],
                is_leader=is_leader
            )
            entry["members_added"] += 1

        if group["members"] and not leader_assigned:
            entry["problems"].append("No Team Leader marked for this project.")

        entries.append(entry)

    return entries


def process_upload(content, supervisor_id):
    """Full pipeline: parse the workbook and apply it. Returns the report."""
    groups, errors = parse_workbook(content)
    entries = apply_upload(groups, supervisor_id) if groups else []

    return {
        "entries": entries,
        "errors": errors,
        "projects_created": sum(1 for e in entries if e["status"] == "created"),
        "projects_updated": sum(1 for e in entries if e["status"] == "updated"),
        "members_added": sum(e["members_added"] for e in entries),
        "accounts_created": sum(len(e["accounts_created"]) for e in entries)
    }
